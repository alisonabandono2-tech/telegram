"""Limpa XLSX localmente para importação. Não usa Telegram nem envia dados."""
import argparse
import re
import unicodedata
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

USERNAME = re.compile(r'^@[A-Za-z][A-Za-z0-9_]{3,31}$')
ALIASES = re.compile(r'(?:rafael|raphael|rafa|rapha)reis|rafinha|raffinha|rafinnha|raphinha')
SERVICES = re.compile(
    r'1win|slots?|slts|cassino|casino|bacbo|tigrinho|br4bet|bra4bet|hrbet|luva\.bet|'
    r'afiliad|affiliat|suporte|support|servicos?|trader|forex|sinais|renda.?extra|'
    r'hotlead|gerente|divulga|promoc|marketing|consultoria|design(?:er)?|assessoria|studio_|\bstudio\b|\bestudio\b|\bloja\b|\bagencia\b|'
    r'https?://|\b\w*bet\b|_bot|\bbot\b|\bganhe\b|\blucre\b|'
    r'win\s+maney|win\s+money|zapamoney|criptoreset|cryptoreset|adicione.?me|'
    r'^grupo\b|^canal\b|^channel\b'
)


def normalize(text):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(text).lower()) if not unicodedata.combining(c))


def clean(frame):
    if not {'Nome','@usuario','ID'}.issubset(frame.columns):
        raise ValueError('A planilha deve conter Nome, @usuario e ID.')
    counts={'total':len(frame),'sem_username_valido':0,'nomes_solicitados':0,'servicos_spam_aparente':0,'duplicatas':0}
    rows=[]
    seen=set()
    for _,row in frame.fillna('').iterrows():
        username=str(row['@usuario']).strip()
        if not USERNAME.fullmatch(username):
            counts['sem_username_valido']+=1
            continue
        name=str(row['Nome']).strip()
        compact=re.sub(r'[^a-z0-9]','',normalize(name+' '+username))
        if ALIASES.search(compact):
            counts['nomes_solicitados']+=1
            continue
        if SERVICES.search(normalize(name+' '+username)) or username.lower().endswith('bot'):
            counts['servicos_spam_aparente']+=1
            continue
        if username.lower() in seen:
            counts['duplicatas']+=1
            continue
        seen.add(username.lower())
        rows.append({'Nome':name,'@usuario':username,'ID':str(row['ID']).strip()})
    counts['mantidos']=len(rows)
    return pd.DataFrame(rows,columns=['Nome','@usuario','ID']),counts


def write_workbook(frame,path):
    frame.to_excel(path,index=False,sheet_name='Contatos',engine='openpyxl')
    book=load_workbook(path)
    sheet=book['Contatos']
    sheet.freeze_panes='A2'
    sheet.auto_filter.ref=sheet.dimensions
    for row in sheet:
        for cell in row:
            if cell.value is not None:
                cell.value=str(cell.value)
                cell.data_type='s'  # texto literal; preserva IDs e impede fórmulas em nomes
    for cell in sheet[1]:
        cell.font=Font(bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='209BD2')
        cell.alignment=Alignment(vertical='center')
    sheet.column_dimensions['A'].width=42
    sheet.column_dimensions['B'].width=32
    sheet.column_dimensions['C'].width=24
    sheet.row_dimensions[1].height=24
    book.save(path)


def main():
    parser=argparse.ArgumentParser(description='Filtra a planilha sem publicar contatos ou enviar mensagens.')
    parser.add_argument('entrada',nargs='?')
    parser.add_argument('saida',nargs='?')
    parser.add_argument('--selecionar',action='store_true',help='Abre seletores de arquivo no Windows')
    args=parser.parse_args()
    root=None
    try:
        if args.selecionar:
            import tkinter as tk
            from tkinter import filedialog,messagebox
            root=tk.Tk();root.withdraw();root.attributes('-topmost',True)
            entrada=filedialog.askopenfilename(parent=root,title='Selecione inscritos_todos.xlsx ORIGINAL',filetypes=[('Planilha Excel','*.xlsx')])
            if not entrada:return
            saida=filedialog.asksaveasfilename(parent=root,title='Salvar planilha filtrada',initialdir=str(Path.home()/'Downloads'),initialfile='inscritos_filtrados.xlsx',defaultextension='.xlsx',filetypes=[('Planilha Excel','*.xlsx')])
            if not saida:return
        else:
            entrada,saida=args.entrada,args.saida
            if not entrada or not saida:parser.error('Informe entrada e saída ou use --selecionar.')
        source,target=Path(entrada),Path(saida)
        if source.resolve()==target.resolve():
            raise ValueError('Escolha outro nome de saída para preservar a planilha original.')
        frame=pd.read_excel(source,dtype=str).fillna('')
        result,counts=clean(frame)
        write_workbook(result,target)
        summary=(f"Planilha pronta: {counts['mantidos']} contatos.\n"
                 f"Sem @ válido: {counts['sem_username_valido']}\n"
                 f"Nomes solicitados: {counts['nomes_solicitados']}\n"
                 f"Serviços/spam aparente: {counts['servicos_spam_aparente']}\n"
                 f"Duplicatas: {counts['duplicatas']}\n\nArquivo: {target}")
        print(summary)
        if root is not None:messagebox.showinfo('Planilha pronta',summary,parent=root)
    except Exception as exc:
        # Não imprime linhas de contatos nem dados da planilha nos erros.
        if root is not None:
            from tkinter import messagebox
            messagebox.showerror('Não foi possível gerar',str(exc) if isinstance(exc,ValueError) else 'Confira o arquivo e feche o Excel antes de tentar novamente.',parent=root)
        raise SystemExit(str(exc) if isinstance(exc,ValueError) else 'Não foi possível processar. Confira o arquivo e feche o Excel.')
    finally:
        if root is not None:root.destroy()


if __name__=='__main__':
    main()
